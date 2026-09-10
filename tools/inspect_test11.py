import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('CleanData/Test11.xlsx', data_only=True)
ws_clean = wb['Cleaned_Data']

# Headers are on row 3
headers_clean = [ws_clean.cell(3, c).value for c in range(1, ws_clean.max_column + 1)]
print('\nCleaned_Data headers (Row 3):')
for idx, h in enumerate(headers_clean, 1):
    print(f'  Col {idx}: {h}')

uni_col_idx = None
id_col_idx = None
phone_col_idx = None
name_col_idx = None
for idx, h in enumerate(headers_clean, 1):
    if h:
        h_str = str(h).strip()
        if 'اسم الجامعة' in h_str or h_str == 'جامعة':
            uni_col_idx = idx
        elif 'الرقم الجامعي' in h_str or 'رقم جامعي' in h_str:
            id_col_idx = idx
        elif 'جوال' in h_str or 'phone' in h_str.lower():
            phone_col_idx = idx
        elif 'اسم' in h_str and 'جامع' not in h_str:
            if not name_col_idx:
                name_col_idx = idx

print(f'\nDetected column indices: Name={name_col_idx}, University={uni_col_idx}, StudentID={id_col_idx}, Phone={phone_col_idx}')

print('\nSample 10 cleaned rows:')
for r in range(4, min(14, ws_clean.max_row + 1)):
    name = ws_clean.cell(r, name_col_idx).value if name_col_idx else 'N/A'
    uni = ws_clean.cell(r, uni_col_idx).value if uni_col_idx else 'N/A'
    st_id = ws_clean.cell(r, id_col_idx).value if id_col_idx else 'N/A'
    ph = ws_clean.cell(r, phone_col_idx).value if phone_col_idx else 'N/A'
    print(f'Row {r}: Name={name} | University={uni} | StudentID={st_id} | Phone={ph}')

# Inspect quarantine
ws_q = wb['quarantine_records']
print('\nQuarantine headers (Row 3):', [ws_q.cell(3, c).value for c in range(1, 7)])

print('\nQuarantined records sample (first 10):')
for r in range(4, min(14, ws_q.max_row + 1)):
    sheet_name = ws_q.cell(r, 2).value
    row_idx = ws_q.cell(r, 3).value
    reason = ws_q.cell(r, 4).value
    payload_str = ws_q.cell(r, 6).value
    print(f'Sheet: {sheet_name}, Row: {row_idx}, Reason: {reason}')
    print(f'  Payload snippet: {str(payload_str)[:200]}')
